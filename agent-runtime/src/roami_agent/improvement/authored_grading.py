"""Independent authored-case checks; student code runs only in disposable sandboxes."""

import base64
import csv
import hashlib
import io
import json
import math
import re
import zipfile

from pptx import Presentation
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from roami_agent.improvement.contracts import DIMENSIONS, digest, relative_file
from roami_agent.improvement.grading import quality
from roami_agent.protocol.errors import DomainError


def close_json(actual, expected, tolerance=1e-5):
    if type(actual) is bool or type(expected) is bool:
        return type(actual) is type(expected) and actual == expected
    if isinstance(expected, dict):
        return (
            isinstance(actual, dict)
            and actual.keys() == expected.keys()
            and all(close_json(actual[k], v, tolerance) for k, v in expected.items())
        )
    if isinstance(expected, list):
        return (
            isinstance(actual, list)
            and len(actual) == len(expected)
            and all(close_json(a, b, tolerance) for a, b in zip(actual, expected, strict=True))
        )
    if type(expected) in {int, float}:
        return (
            type(actual) in {int, float}
            and math.isfinite(actual)
            and abs(actual - expected) <= tolerance * max(1, abs(expected))
        )
    return actual == expected


def corrected_backtest_matches(actual, expected, prices):
    if not isinstance(actual, dict) or not isinstance(actual.get("strategy"), dict):
        return False
    strategy = actual["strategy"]
    reference = expected["strategy"]
    cash, quantity = prices["initial_cash"], 0
    for trade in reference["trades"]:
        sign = 1 if trade["side"] == "buy" else -1
        quantity += sign * trade["quantity"]
        cash -= sign * trade["quantity"] * trade["price"] + trade["fee"]
    summaries = {
        "initial_cash": prices["initial_cash"],
        "fee": prices["fee"],
        "slippage": prices["slippage"],
        "final_cash": cash,
        "final_quantity": quantity,
        "final_equity": reference["equity"][-1],
    }
    extras = strategy.keys() - reference.keys()
    if not extras <= summaries.keys() | {"description"}:
        return False
    for key in extras:
        if key == "description":
            if not isinstance(strategy[key], str) or not strategy[key].strip():
                return False
        elif not close_json(strategy[key], summaries[key]):
            return False
    normalized = {**actual, "strategy": {k: v for k, v in strategy.items() if k not in extras}}
    return close_json(normalized, expected)


def read_file(workspace, name, *, binary=False):
    path = relative_file(workspace, name)
    if not path.is_file() or path.stat().st_size > 16_000_000:
        raise ValueError("artifact unavailable")
    return path.read_bytes() if binary else path.read_text()


async def grade_authored(case, workspace, executor):
    spec = case["judge"]
    checks = {}
    evidence = []
    try:
        for name in spec["required"]:
            value = read_file(workspace, name, binary=True)
            checks["artifact:" + name] = bool(value)
            evidence.append({"path": name, "sha256": hashlib.sha256(value).hexdigest()})
        for name, expected_hash in spec.get("protected", {}).items():
            checks["protected:" + name] = digest(read_file(workspace, name)) == expected_hash
        for name, expected_hash in spec.get("protected_assets", {}).items():
            checks["protected:" + name] = (
                hashlib.sha256(read_file(workspace, name, binary=True)).hexdigest() == expected_hash
            )
        if not all(checks.values()):
            raise ValueError("protected input changed")
        identifier = spec["case_id"]
        if "expected" in spec:
            actual = json.loads(read_file(workspace, "result.json"))
            compared = actual
            if identifier == "STOCK-03":
                scenario_fields = {"scenarios", "falsification"}
                compared = (
                    {k: v for k, v in actual.items() if k not in scenario_fields}
                    if isinstance(actual, dict)
                    else actual
                )
                scenarios = actual.get("scenarios") if isinstance(actual, dict) else None
                falsification = actual.get("falsification") if isinstance(actual, dict) else None
                checks["conditional_scenarios"] = (
                    isinstance(scenarios, list)
                    and len(scenarios) >= 2
                    and all(
                        isinstance(s, str)
                        and bool(s.strip())
                        or isinstance(s, dict)
                        and isinstance(s.get("condition"), str)
                        and bool(s["condition"].strip())
                        for s in scenarios
                    )
                )
                checks["falsification_conditions"] = (
                    isinstance(falsification, str)
                    and bool(falsification.strip())
                    or isinstance(falsification, list)
                    and bool(falsification)
                    and all(isinstance(s, str) and bool(s.strip()) for s in falsification)
                )
            checks["structured_result"] = close_json(
                compared, spec["expected"], spec.get("numeric_tolerance", 1e-5)
            )
        if "report.md" in spec["required"]:
            report = read_file(workspace, "report.md")
            receipt = case["fixture_receipt"]
            checks["report_grounding"] = len(report) >= 40 and receipt in report
            for text in spec.get("text_required", []):
                checks["report:" + text] = text in report
        if identifier == "CODE-01":
            output = json.loads(
                (await executor.sandbox(workspace, API_PROBE)).decode().splitlines()[-1]
            )
            expected_items = sorted(spec["items"], key=lambda i: i["name"])
            responses = output["responses"]
            if spec.get("pagination_metadata"):
                responses = [
                    [r[0], {key: r[1].get(key) for key in ("items", "total")}]
                    if isinstance(r[1], dict) and r[0] == 200
                    else r
                    for r in responses
                ]
            checks["default_page"] = responses[0] == [
                200,
                {"items": expected_items[:5], "total": 12},
            ]
            checks["third_page"] = responses[1] == [
                200,
                {"items": expected_items[10:], "total": 12},
            ]
            checks["outside_page"] = responses[2] == [200, {"items": [], "total": 12}]
            checks["invalid_parameters"] = all(r[0] == 400 for r in responses[3:7])
            checks["health_regression"] = responses[7][0] == 200
            checks["actual_student_tests"] = output["tests_passed"] and output["test_count"] >= 3
        elif identifier == "OPS-04":
            output = json.loads(
                (await executor.sandbox(workspace, BACKUP_PROBE)).decode().splitlines()[-1]
            )
            checks["restore_roundtrip"] = output["restored"] == spec["source"]
            checks["reject_incompatible"] = output["incompatible_rejected"]
            checks["restore_unchanged_on_error"] = output["target_unchanged"]
        elif identifier == "DOC-03":
            source = base64.b64decode(case["asset_files"]["source.pptx"])
            original = Presentation(io.BytesIO(source))
            updated = Presentation(str(workspace / "updated.pptx"))
            checks["editable_pages"] = len(updated.slides) == len(original.slides) == 6
            texts = [
                "\n".join(s.text for s in slide.shapes if s.has_text_frame)
                for slide in updated.slides
            ]
            checks["targeted_text"] = (
                "135" in texts[2]
                and "120" not in texts[2]
                and "8日" in texts[4]
                and "5日" not in texts[4]
            )
            with (
                zipfile.ZipFile(io.BytesIO(source)) as a,
                zipfile.ZipFile(workspace / "updated.pptx") as b,
            ):
                checks["untouched_pages"] = all(
                    a.read(f"ppt/slides/slide{i}.xml") == b.read(f"ppt/slides/slide{i}.xml")
                    for i in [1, 2, 4, 6]
                )
                media = [n for n in a.namelist() if n.startswith("ppt/media/")]
                checks["resources_preserved"] = all(
                    n in b.namelist() and a.read(n) == b.read(n) for n in media
                )
            preview = PdfReader(io.BytesIO(read_file(workspace, "preview.pdf", binary=True)))
            checks["rendered_preview"] = len(preview.pages) == 6
            checks["preview_content"] = (
                len(preview.pages) == 6
                and "135" in preview.pages[2].extract_text()
                and "8" in preview.pages[4].extract_text()
            )
        elif identifier == "DOC-05":
            data = json.loads(read_file(workspace, "result.json"))
            expected = spec["expected_constraints"]
            checks["constraints"] = all(data.get(k) == v for k, v in expected.items() if k != "new")
            ids = {"project_path", "budget", "immutable_api", "mobile", "offline_recovery"}
            checks["stories_acceptance"] = (
                isinstance(data.get("requirements"), list)
                and all(isinstance(v, str) for v in data["requirements"])
                and set(data["requirements"]) == ids
                and isinstance(data.get("stories"), dict)
                and set(data["stories"]) == ids
                and isinstance(data.get("acceptance"), dict)
                and set(data["acceptance"]) == ids
                and data.get("interfaces") == ["/v1/report"]
            )
            checks["architecture"] = "<svg" in read_file(workspace, "architecture.svg") and all(
                i in read_file(workspace, "plan.md") for i in ids
            )
        elif identifier.startswith("DEMO-"):
            program = (
                BOARD_PROBE
                if identifier == "DEMO-01"
                else RESPONSIVE_PROBE
                if identifier == "DEMO-05"
                else CSV_PROBE
            )
            output = json.loads(
                (await executor.sandbox(workspace, program, timeout=90)).decode().splitlines()[-1]
            )
            checks.update({"browser:" + k: bool(v) for k, v in output.items()})
        elif identifier.startswith("QUANT-"):
            prices = json.loads(read_file(workspace, "prices.json"))
            output = json.loads(
                (await executor.sandbox(workspace, RERUN_ANALYSIS, timeout=60))
                .decode()
                .splitlines()[-1]
            )
            checks["reproducible_code"] = (
                corrected_backtest_matches(output["corrected"], spec["corrected_expected"], prices)
                if identifier == "QUANT-03"
                else close_json(output["result"], spec["expected"])
            )
            if identifier == "QUANT-01":
                trades = list(csv.DictReader(read_file(workspace, "trades.csv").splitlines()))
                normalized = [
                    {k: (v if k == "side" else float(v)) for k, v in row.items()} for row in trades
                ]
                checks["transaction_file"] = close_json(normalized, spec["expected"]["trades"])
                checks["curve"] = "<svg" in read_file(workspace, "equity.svg")
            if identifier == "QUANT-02":
                log = json.loads(read_file(workspace, "selection-log.json"))
                checks["selection_log"] = (
                    log.get("candidates") == spec["expected"]["tried"]
                    and log.get("selection_range") == [100, 140]
                    and log.get("test_used_for_selection") is False
                )
            if identifier == "QUANT-03":
                checks["corrected_run_verified"] = corrected_backtest_matches(
                    json.loads(read_file(workspace, "corrected.json")),
                    spec["corrected_expected"],
                    prices,
                ) and corrected_backtest_matches(
                    output["corrected"], spec["corrected_expected"], prices
                )
                code = read_file(workspace, "analysis.py")
                checks["no_original_future_signal"] = not re.search(r"\.shift\s*\(\s*-1\s*\)", code)
        elif identifier == "LAB-05":
            experience = json.loads(read_file(workspace, "experience.json"))
            log = json.loads(read_file(workspace, "transfer-log.json"))
            checks["sourced_experience"] = (
                experience.get("source") == "project/README.md"
                and experience.get("entry") == "python verify_current.py"
            )
            outputs = [log[key] for key in ("output", "stdout") if key in log]
            checks["actual_entry_output"] = bool(outputs) and all(
                isinstance(value, str) and value.strip() == "VERIFIED:" + case["fixture_receipt"]
                for value in outputs
            )
        if identifier.startswith("NEWS-"):
            json.loads(read_file(workspace, "sources.json"))
            report = read_file(workspace, "report.md")
            expected = spec["expected"]
            urls = (
                [i["source"] for i in expected["items"]]
                if "items" in expected
                else [expected.get("source", expected.get("confirmed_source", ""))]
            )
            checks["cited_sources"] = all(url and url in report for url in urls)
    except (
        ValueError,
        OSError,
        KeyError,
        TypeError,
        DomainError,
        PdfReadError,
        zipfile.BadZipFile,
    ) as error:
        checks["verification_completed"] = False
        failure = {"verification_error": type(error).__name__}
        if isinstance(error, DomainError) and error.code in {
            "RSI_DOCKER_COMMAND_FAILED",
            "RSI_DOCKER_OUTPUT_LIMIT",
            "RSI_TIME_BUDGET",
            "RSI_ARTIFACT_UNAVAILABLE",
        }:
            failure["verification_error_code"] = error.code
        evidence.append(failure)
    hard_passed = bool(checks) and all(checks.values())
    pending = hard_passed and spec.get("rubric", False)
    grade = quality(
        dict.fromkeys(DIMENSIONS, 3 if hard_passed else 0),
        case.get("grading_profile", "mechanism"),
        hard_passed=hard_passed,
        pending=pending,
    )
    if pending:
        grade["quality"] = 0
    grade.update(
        checks=checks,
        evidence=evidence,
        artifact_hash=digest(evidence),
        grading_profile=case.get("grading_profile", "mechanism"),
        rubric_required=spec.get("rubric", False),
        reason="independent_rubric_required"
        if pending
        else "objective_checks_passed"
        if hard_passed
        else "hard_assertion_failed",
        grade_scope="authored_contract",
    )
    return grade


API_PROBE = r"""import os,sys,time,json,subprocess,urllib.request,urllib.error,re,socket
s=socket.socket();s.bind(("127.0.0.1",0));port=s.getsockname()[1];s.close()
p=subprocess.Popen([sys.executable,"app.py"],env={**os.environ,"PORT":str(port)},stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
try:
 for _ in range(100):
  try: urllib.request.urlopen(f"http://127.0.0.1:{port}/health",timeout=.2);break
  except Exception: time.sleep(.05)
 results=[]
 for path in ["/items","/items?page=3","/items?page=99","/items?page=0","/items?page=x","/items?pageSize=0","/items?pageSize=101","/health"]:
  try:
   r=urllib.request.urlopen(f"http://127.0.0.1:{port}"+path);results.append([r.status,json.load(r)])
  except urllib.error.HTTPError as e:results.append([e.code,None])
 result=subprocess.run([sys.executable,"-m","unittest","discover","-p","test_items.py"],capture_output=True,text=True,timeout=15)
 match=re.search(r"Ran (\d+) tests?",result.stderr)
 print(json.dumps({"responses":results,"tests_passed":result.returncode==0,"test_count":int(match[1]) if match else 0}))
finally:p.terminate();p.wait()
"""

BACKUP_PROBE = r"""import tempfile,pathlib,subprocess,sys,json,zipfile
with tempfile.TemporaryDirectory() as d:
 root=pathlib.Path(d);archive=root/"backup.zip";target=root/"target"
 subprocess.run([sys.executable,"backup_restore.py","backup","source",str(archive)],check=True,capture_output=True)
 subprocess.run([sys.executable,"backup_restore.py","restore",str(archive),str(target)],check=True,capture_output=True)
 restored={str(p.relative_to(target)):p.read_text() for p in target.rglob("*") if p.is_file()}
 bad=root/"bad.zip"
 with zipfile.ZipFile(archive) as a,zipfile.ZipFile(bad,"w") as b:
  for name in a.namelist():
   raw=a.read(name)
   if name.endswith("database.json"):data=json.loads(raw);data["schema"]=99;raw=json.dumps(data).encode()
   b.writestr(name,raw)
 before={str(p.relative_to(target)):p.read_bytes().hex() for p in target.rglob("*") if p.is_file()}
 result=subprocess.run([sys.executable,"backup_restore.py","restore",str(bad),str(target)],capture_output=True)
 after={str(p.relative_to(target)):p.read_bytes().hex() for p in target.rglob("*") if p.is_file()}
 print(json.dumps({"restored":restored,"incompatible_rejected":result.returncode!=0,"target_unchanged":before==after}))
"""

RERUN_ANALYSIS = r"""import tempfile,pathlib,shutil,subprocess,sys,json
with tempfile.TemporaryDirectory() as d:
 root=pathlib.Path(d)
 for path in pathlib.Path("/workspace").iterdir():
  if path.is_file():shutil.copy2(path,root/path.name)
 for name in ["result.json","corrected.json","trades.csv","equity.svg","selection-log.json"]:(root/name).unlink(missing_ok=True)
 subprocess.run([sys.executable,"analysis.py"],cwd=root,check=True,capture_output=True,timeout=30)
 print(json.dumps({"result":json.loads((root/"result.json").read_text()) if (root/"result.json").exists() else None,"corrected":json.loads((root/"corrected.json").read_text()) if (root/"corrected.json").exists() else None}))
"""

BOARD_PROBE = r"""import json,pathlib
from playwright.sync_api import sync_playwright
with sync_playwright() as p:
 b=p.chromium.launch(executable_path="/usr/bin/chromium",headless=True,args=["--no-sandbox"]);page=b.new_page(viewport={"width":1280,"height":900});page.set_default_timeout(5000);page.goto("file:///workspace/index.html");page.wait_for_timeout(200)
 before=page.get_by_test_id("board-card").count();page.get_by_test_id("new-title").fill("RSI-added");page.get_by_test_id("add-card").click();added=page.get_by_test_id("board-card").count()==before+1
 card=page.get_by_test_id("board-card").filter(has_text="RSI-added");card.drag_to(page.locator('[data-status="done"]:not([data-testid="board-card"])').first)
 moved=card.get_attribute("data-status")=="done";page.reload();page.wait_for_timeout(100);persisted=page.get_by_test_id("board-card").filter(has_text="RSI-added").get_attribute("data-status")=="done"
 print(json.dumps({"initial_five":before==5,"added":added,"moved":moved,"persisted":persisted}));b.close()
"""

RESPONSIVE_PROBE = r"""import json
from playwright.sync_api import sync_playwright
with sync_playwright() as p:
 b=p.chromium.launch(executable_path="/usr/bin/chromium",headless=True,args=["--no-sandbox"]);page=b.new_page();page.set_default_timeout(5000);page.goto("file:///workspace/index.html");page.get_by_test_id("draft").fill("draft-preserved");page.get_by_test_id("filter").fill("Card 1")
 fits=True;touch=True
 for width in [320,390,820]:
  page.set_viewport_size({"width":width,"height":800});fits &= page.evaluate("document.documentElement.scrollWidth<=innerWidth");touch &= page.get_by_test_id("draft").bounding_box()["height"]>=44
 page.reload();page.wait_for_timeout(100)
 print(json.dumps({"responsive":fits,"touch":touch,"draft":page.get_by_test_id("draft").input_value()=="draft-preserved","filter":page.get_by_test_id("filter").input_value()=="Card 1","results":page.get_by_test_id("board-card").count()==1}));b.close()
"""

CSV_PROBE = r"""import json,os,socket,subprocess,sys,time,urllib.request
from playwright.sync_api import sync_playwright
s=socket.socket();s.bind(("127.0.0.1",0));port=s.getsockname()[1];s.close();proc=subprocess.Popen([sys.executable,"server.py"],env={**os.environ,"PORT":str(port)},stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
try:
 for _ in range(100):
  try:urllib.request.urlopen(f"http://127.0.0.1:{port}/stats");break
  except Exception:time.sleep(.05)
 with sync_playwright() as p:
  b=p.chromium.launch(executable_path="/usr/bin/chromium",headless=True,args=["--no-sandbox"]);page=b.new_page();page.set_default_timeout(5000);page.goto(f"http://127.0.0.1:{port}");file=page.get_by_test_id("csv-input");button=page.get_by_test_id("import")
  def submit():
   with page.expect_response(lambda response: response.request.method=="POST" and response.url==f"http://127.0.0.1:{port}/import") as response:button.click()
   response.value.finished()
  file.set_input_files("/workspace/valid.csv");submit();first=json.load(urllib.request.urlopen(f"http://127.0.0.1:{port}/stats"))["rows"]==2
  file.set_input_files("/workspace/invalid.csv")
  page.wait_for_function("() => !document.querySelector('[data-testid=import]').disabled || Boolean(document.querySelector('[data-testid=error]').textContent.trim())")
  if button.is_enabled():button.click()
  page.wait_for_function("() => Boolean(document.querySelector('[data-testid=error]').textContent.trim())")
  missing=bool(page.get_by_test_id("error").inner_text())
  file.set_input_files("/workspace/valid.csv");submit()
  page.wait_for_function("() => Boolean(document.querySelector('[data-testid=error]').textContent.trim())")
  preserved=file.evaluate("e=>e.files.length===1") and bool(page.get_by_test_id("error").inner_text())
  if preserved:submit()
  stats=json.load(urllib.request.urlopen(f"http://127.0.0.1:{port}/stats"));print(json.dumps({"valid_import":first,"missing_column":missing,"input_preserved":preserved,"exactly_once_retry":preserved and stats["rows"]==4 and stats["attempts"]==3}));b.close()
finally:proc.terminate();proc.wait()
"""
