"""Original stock data equality checkpoint with native upload evidence; quality stays pending."""

import ast
import base64
import hashlib
import json
import re

from roami_agent.improvement.domain.contracts import DIMENSIONS, digest, require
from roami_agent.improvement.domain.scoring import quality
from roami_agent.improvement.evaluation.contracts import freeze_suite
from roami_agent.improvement.evaluation.execution.stock_environment import IMAGE_RE, OUTPUTS
from roami_agent.improvement.integrations.artifact_paths import relative_file
from roami_agent.protocol.errors import DomainError

TASK_ID = "ds-stock-analysis-slides"


def import_stock(source_dir, *, image, raw_source, revision):
    require(isinstance(image, str) and re.fullmatch(IMAGE_RE, image))
    require(bool(re.fullmatch(r"[a-f0-9]{40}", revision)))
    task = source_dir / "workspaces" / "tasks" / TASK_ID
    instruction = (task / "task.md").read_text()
    verifier = (task / "evaluator.py").read_text()
    golden = (task / "stock_data.csv").read_bytes()
    raw = raw_source.read_bytes()
    require(0 < len(raw) <= 2_000_000 and 0 < len(golden) <= 2_000_000)
    node = next(
        n
        for n in ast.parse(verifier).body
        if isinstance(n, ast.FunctionDef) and n.name == "grade_checkpoint1"
    )
    checkpoint = ast.get_source_segment(verifier, node)
    require(checkpoint and "golden_df" in checkpoint)
    case = {
        "id": "the-agent-company:" + TASK_ID,
        "category": "LIFE",
        "layer": "public_component",
        "split": "public_test",
        "cluster": "the-agent-company:stock-workflow",
        "counts_as_delivery": True,
        "grading_profile": "presentation",
        "files": {},
        "prompt": instruction
        + "\n\nUse the isolated native ownCloud through python /workspace/stock-http.py GET|MKCOL|PUT /Documents/path [local-file]. GET downloads to local-file; PUT uploads it. Only the supplied raw CSV and the three requested output filenames are accessible, and MKCOL creates the requested stock_presentation folder. The agent has no direct network or service credentials. Native state is captured after exit: local files alone do not count as delivery. Do not replay writes with unknown outcomes; reconcile by reading.",
        "limits": {"seconds": 900, "tokens": 256000, "requests": 48},
        "judge": {
            "kind": "company_stock",
            "task_id": TASK_ID,
            "image": image,
            "raw_sha256": hashlib.sha256(raw).hexdigest(),
            "golden": base64.b64encode(golden).decode(),
            "golden_sha256": hashlib.sha256(golden).hexdigest(),
            "checkpoint": checkpoint,
            "verifier": verifier,
            "rubric_required": True,
        },
        "origin": {
            "benchmark": "the_agent_company",
            "task_id": TASK_ID,
            "revision": revision,
            "source_split": "company_stock_workflow",
            "adapter_revision": "native-owncloud-post-exit-webdav-v2",
            "input_sha256": digest(
                {
                    "instruction": instruction,
                    "raw_sha256": hashlib.sha256(raw).hexdigest(),
                    "image": image,
                }
            ),
            "verifier_sha256": hashlib.sha256(verifier.encode()).hexdigest(),
            "adapted": True,
            "adaptation": "Publisher ownCloud1.0.0 seeded image pinned by digest, fresh anonymous data volume per trial, non-root internal unpublished service, allowlisted Unix-socket WebDAV. Original CSV equality checkpoint retained; original seven LLM chart/script points unscored pending independent calibration. Adapted numeric gate rtol=1e-12/atol=1e-10 handles seven published float roundoff artifacts; the original strict equality result remains separately reported. Extra source-immutability and remote-delivery structural gates. Agent instructions contain only the original task and transport contract; grader tolerances and reference ordering are not data-cleaning instructions. No browser or official full score.",
        },
    }
    return freeze_suite(
        "the-agent-company-stock-native-v2", [case], official_leaderboard_comparable=False
    )


PRESENTATION_CHECK = r"""from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
def plot_slide_count(path):
 def has_plot(shapes):
  for shape in shapes:
   if shape.has_chart or getattr(shape,'image',None) is not None:return True
   if shape.shape_type==MSO_SHAPE_TYPE.GROUP and has_plot(shape.shapes):return True
  return False
 return sum(has_plot(slide.shapes) for slide in Presentation(path).slides)
"""


PROGRAM = (
    PRESENTATION_CHECK
    + r"""import base64,io,json,pathlib,sys
import pandas as pd
spec=json.loads(sys.stdin.read());root=pathlib.Path('/workspace');golden=base64.b64decode(spec['golden']);read_csv=pd.read_csv
pd.read_csv=lambda source,*a,**kw:read_csv(io.BytesIO(golden) if source=='/utils/stock_data.csv' else source,*a,**kw)
namespace={'pd':pd,'BytesIO':io.BytesIO,'logging':__import__('logging'),'check_file_in_owncloud_directory':lambda f,d:(root/f).is_file(),'get_binary_file_content_owncloud':lambda f,d:(root/f).read_bytes()}
exec(compile(spec['checkpoint'],'original_csv_checkpoint.py','exec'),namespace)
checks={};native_csv_equality=None
try:native_csv_equality=bool(namespace['grade_checkpoint1']())
except Exception:checks['native_checkpoint_completed']=False
try:
 pd.testing.assert_frame_equal(read_csv(root/'stock_cleaned.csv'),read_csv(io.BytesIO(golden)),check_exact=False,rtol=1e-12,atol=1e-10)
 checks['adapted_csv_values']=True
except AssertionError:checks['adapted_csv_values']=False
except (pd.errors.ParserError,pd.errors.EmptyDataError,UnicodeError):
 checks['adapted_csv_values']=False;checks['csv_readable']=False
except Exception:
 checks['adapted_csv_values']=False;checks['csv_verification_completed']=False
try:
 checks['five_plot_slides']=plot_slide_count(root/'stock_analysis.pptx')>=5
except Exception:checks['five_plot_slides']=False
try:checks['nonempty_utf8_script']=len((root/'stock_analysis_script.md').read_text().strip())>=100
except (OSError,UnicodeError):checks['nonempty_utf8_script']=False
print(json.dumps({'checks':checks,'native_csv_equality':native_csv_equality}))"""
)


async def grade_stock(case, executor, *, root):
    checks, evidence, native_csv = {}, [], None
    try:
        state_path = relative_file(root, "stock-snapshot.json")
        require(state_path.is_file() and state_path.stat().st_size < 100000)
        state = json.loads(state_path.read_text())
        checks["source_unchanged"] = (
            state["initial"]["source_sha256"]
            == state["final"]["source_sha256"]
            == case["judge"]["raw_sha256"]
        )
        checks["fresh_remote_outputs"] = not state["initial"]["outputs"] and set(
            state["final"]["outputs"]
        ) == set(OUTPUTS)
        remote = relative_file(root, "stock-remote")
        for name in OUTPUTS:
            file = relative_file(remote, name)
            require(file.is_file() and file.stat().st_size <= 8_000_000)
            sha = hashlib.sha256(file.read_bytes()).hexdigest()
            require(sha == state["final"]["outputs"][name]["sha256"], "RSI_SOURCE_CHANGED")
            evidence.append({"remote_path": name, "sha256": sha})
        result = json.loads(
            (
                await executor.sandbox(
                    remote, PROGRAM, input=json.dumps(case["judge"]).encode(), timeout=90
                )
            )
            .decode()
            .splitlines()[-1]
        )
        checks.update(result["checks"])
        native_csv = result["native_csv_equality"]
        evidence.append(
            {
                "source": "host_native_snapshot",
                "sha256": hashlib.sha256(state_path.read_bytes()).hexdigest(),
            }
        )
    except (DomainError, OSError, ValueError, KeyError, TypeError):
        checks["verification_completed"] = False
    hard = bool(checks) and all(checks.values())
    grade = quality(
        dict.fromkeys(DIMENSIONS, 3 if hard else 0), "presentation", hard_passed=hard, pending=hard
    )
    grade.update(
        quality=0,
        checks=checks,
        evidence=evidence,
        artifact_hash=digest(evidence),
        native_partial={
            "scored_points": int(native_csv) if native_csv is not None else None,
            "total_points": 8,
            "unscored_points": 7 if native_csv is not None else 8,
        },
        rubric_required=True,
        grade_scope="native_uploaded_stock_artifacts_csv_checkpoint_only_quality_pending",
        reason="independent_rubric_required" if hard else "hard_assertion_failed",
    )
    return grade
