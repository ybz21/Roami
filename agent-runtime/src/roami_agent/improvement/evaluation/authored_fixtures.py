"""Authored workload inputs and independent oracles; excluded from agent images."""

import base64
import hashlib
import io
import json
import math
from datetime import datetime, timedelta

from pptx import Presentation
from pptx.util import Inches, Pt

from roami_agent.improvement.domain.contracts import digest


def prices(count=60):
    rows = []
    for i in range(count):
        close = round(100 + 8 * math.sin(i / 3) + i / 8, 2)
        rows.append({"bar": i, "open": round(close + (3 if i % 9 == 0 else -1), 2), "close": close})
    return rows


def backtest(rows, fast=3, slow=5, fee=0.001, slip=0.001):
    cash, quantity, pending = 10000.0, 0.0, None
    trades, equity = [], []
    closes = []
    for i, bar in enumerate(rows):
        if pending is not None:
            if pending and not quantity:
                price = bar["open"] * (1 + slip)
                quantity = cash / (price * (1 + fee))
                cost = quantity * price * fee
                trades.append(
                    {
                        "signal_bar": i - 1,
                        "execution_bar": i,
                        "side": "buy",
                        "price": price,
                        "quantity": quantity,
                        "fee": cost,
                    }
                )
                cash = 0
            elif not pending and quantity:
                price = bar["open"] * (1 - slip)
                cost = quantity * price * fee
                trades.append(
                    {
                        "signal_bar": i - 1,
                        "execution_bar": i,
                        "side": "sell",
                        "price": price,
                        "quantity": quantity,
                        "fee": cost,
                    }
                )
                cash = quantity * price - cost
                quantity = 0
        closes.append(bar["close"])
        pending = (
            sum(closes[-fast:]) / fast > sum(closes[-slow:]) / slow if len(closes) >= slow else None
        )
        equity.append(cash + quantity * bar["close"])
    return {"trades": trades, "equity": equity}


def presentation():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for i in range(6):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(2))
        shape.text = {2: "指标 120", 4: "计划 5日"}.get(i, "已确认页面 " + str(i + 1))
        for p in shape.text_frame.paragraphs:
            p.font.size = Pt(36)
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def news_sources(nonce):
    stamp = datetime.fromisoformat("2026-10-22T09:00:00+08:00")
    sources = [
        {
            "event_id": "e" + str(i),
            "source": f"https://fixture.invalid/official/{i}",
            "published_at": (stamp - timedelta(hours=i)).isoformat(),
            "occurred_at": (stamp - timedelta(hours=i + 1)).isoformat(),
            "title": f"虚构技术事件{i}",
            "evidence": nonce + str(i),
            "kind": "official",
        }
        for i in range(1, 6)
    ]
    return sources + [
        {**sources[0], "source": "https://fixture.invalid/reprint/a", "kind": "reprint"},
        {**sources[1], "source": "https://fixture.invalid/reprint/b", "kind": "reprint"},
        {
            **sources[2],
            "event_id": "old",
            "published_at": (stamp - timedelta(hours=30)).isoformat(),
        },
    ]


def recipe(case_id, seed):
    nonce = hashlib.sha256((seed + case_id).encode()).hexdigest()[:24]
    files, assets, protected = {}, {}, []
    judge = {"kind": "authored", "case_id": case_id}
    required = ["result.json"]
    schema = {}
    if case_id == "CODE-01":
        rows = [
            {"id": i, "name": name}
            for i, name in enumerate(
                ["L", "A", "K", "B", "J", "C", "I", "D", "H", "E", "G", "F"], 1
            )
        ]
        files["items.json"] = json.dumps(rows)
        files["app.py"] = """from http.server import BaseHTTPRequestHandler, HTTPServer
import json,os
class Handler(BaseHTTPRequestHandler):
 def do_GET(self):
  self.send_response(200 if self.path=="/health" else 404);self.end_headers();self.wfile.write(b'{"ok":true}')
if __name__=="__main__": HTTPServer(("127.0.0.1",int(os.getenv("PORT","8000"))),Handler).serve_forever()
"""
        required = ["app.py", "test_items.py"]
        schema = {
            "contract": "app.py 从PORT读端口；/health保持200；非法page/pageSize返回400，pageSize上限100；响应必要字段为items与total，可附带page/pageSize元数据；补unittest测试。"
        }
        judge["items"] = rows
        judge["pagination_metadata"] = True
    elif case_id == "OPS-04":
        rows = [{"id": i, "value": nonce + str(i)} for i in range(100)]
        files["source/database.json"] = json.dumps({"schema": 1, "records": rows})
        for i in range(3):
            files[f"source/attachment-{i}.txt"] = nonce + "-" + str(i)
        files["source/hashes.json"] = json.dumps(
            {
                k.removeprefix("source/"): hashlib.sha256(v.encode()).hexdigest()
                for k, v in files.items()
            }
        )
        required = ["backup_restore.py", "backup.zip", "restored/database.json", "result.json"]
        schema = {
            "contract": "backup_restore.py CLI: backup SOURCE ZIP / restore ZIP TARGET。保留source原始输入，只删除演练副本；还原目标含完整source内容；schema!=1拒绝且保持已有目标。生成可核对result.json和真实backup.zip。"
        }
        judge["source"] = {k.removeprefix("source/"): v for k, v in files.items()}
    elif case_id == "DOC-03":
        assets["source.pptx"] = base64.b64encode(presentation()).decode()
        required = ["updated.pptx", "changes.md", "preview.pdf"]
        judge["rubric"] = True
        schema = {
            "contract": "第3页指标120改为135，第5页计划5日改为8日。编辑source.pptx，保存updated.pptx；生成changes.md与LibreOffice导出的preview.pdf。源文件保持；未改页XML语义、文本与对象保持。"
        }
    elif case_id == "DOC-05":
        inputs = {
            "project_path": "/workspace/" + nonce,
            "budget": 5000,
            "immutable_api": "/v1/report",
            "new": ["mobile", "offline_recovery"],
            "source_receipt": nonce,
        }
        files["previous-plan.json"] = json.dumps(inputs)
        required = ["plan.md", "architecture.svg", "result.json"]
        schema = {
            "contract": "result.json顶层保留previous-plan.json的project_path、budget、immutable_api、source_receipt，字段值和类型不变。requirements为需求ID字符串数组；stories和acceptance为需求ID到文字的对象；interfaces为字符串数组。",
            "requirements": [
                "project_path",
                "budget",
                "immutable_api",
                "mobile",
                "offline_recovery",
            ],
            "stories": "每项需求一个故事",
            "interfaces": ["/v1/report"],
            "acceptance": "逐项关联需求ID",
            "source_receipt": "原始receipt",
        }
        judge.update(expected_constraints=inputs, rubric=True)
    elif case_id.startswith("DEMO-"):
        files["index.html"] = (
            '<!doctype html><meta charset="utf-8"><title>Demo</title><main id="app">Implement requested demo.</main>'
        )
        files["initial.json"] = json.dumps(
            {
                "cards": [
                    {
                        "id": str(i),
                        "title": "Card " + str(i),
                        "status": ["todo", "doing", "done"][i % 3],
                    }
                    for i in range(5)
                ],
                "receipt": nonce,
            }
        )
        required = ["index.html", "README.md"]
        if case_id == "DEMO-01":
            schema = {
                "contract": "data-testid: new-title输入、add-card按钮、board-card卡片（data-id、data-status）、列data-status=todo/doing/done。卡片可拖动，localStorage持久化；初始5卡。index.html直接file://打开。"
            }
        elif case_id == "DEMO-05":
            files["existing.js"] = (
                "// 已确认契约：筛选Card标题；草稿与筛选使用localStorage持久化。\n"
            )
            schema = {
                "contract": "data-testid: filter、draft、board-card；加载initial.json对应5卡，筛选Card标题。输入草稿、筛选在刷新与320/390/820宽度保持；控件>=44px；不横向溢出。可将初始JSON内联；file://入口。"
            }
        else:
            files["valid.csv"] = "name,amount\nAlpha,17\nBeta,23\n"
            files["invalid.csv"] = "name\nAlpha\n"
            files["server.py"] = CSV_SERVER
            protected.append("server.py")
            schema = {
                "contract": "服务由grader以PORT启动server.py，禁止修改。UI data-testid: csv-input文件、import按钮、error错误、rows结果列表。先本地检查name/amount列；POST /import JSON {rows:[{name,amount}]}携带Idempotency-Key。重新选择文件开始一次新导入，即使文件内容与已成功导入的内容相同也用新键；同一次导入遇503后保留文件、内容与原键，直接点击重试沿用原键，不重复写入。不按文件内容永久去重。GET /stats提供真实统计。"
            }
        judge.update(rubric=True)
    elif case_id.startswith("NEWS-"):
        required = ["report.md", "result.json"]
        if case_id == "NEWS-01":
            sources = news_sources(nonce)
            files["sources.json"] = json.dumps(
                {"as_of": "2026-10-22T09:00:00+08:00", "sources": sources}, ensure_ascii=False
            )
            expected = {
                "items": [
                    {
                        k: s[k]
                        for k in ("event_id", "source", "published_at", "occurred_at", "evidence")
                    }
                    for s in sources[:5]
                ]
            }
        elif case_id == "NEWS-02":
            sources = [
                {
                    "source": "https://fixture.invalid/official/old",
                    "date": "2026-10-22",
                    "venue": "A馆",
                    "receipt": nonce + "1",
                },
                {
                    "source": "https://fixture.invalid/official/correction",
                    "date": "2026-10-23",
                    "venue": "B馆",
                    "replaces": "https://fixture.invalid/official/old",
                    "receipt": nonce + "2",
                },
            ]
            files["sources.json"] = json.dumps(sources, ensure_ascii=False)
            expected = {
                "date": "2026-10-23",
                "venue": "B馆",
                "source": sources[1]["source"],
                "replaces": sources[0]["source"],
                "receipt": nonce + "2",
            }
        else:
            files["sources.json"] = json.dumps(
                {
                    "official": {
                        "launch": "2026-10-23",
                        "parameter": None,
                        "source": "https://fixture.invalid/official",
                        "receipt": nonce,
                    },
                    "reprints": [
                        {"origin": "rumor-1", "parameter": 42},
                        {"origin": "rumor-1", "parameter": 42},
                    ],
                }
            )
            expected = {
                "launch": "2026-10-23",
                "parameter": None,
                "confirmed_source": "https://fixture.invalid/official",
                "independent_rumor_sources": 1,
                "receipt": nonce,
            }
        judge["expected"] = expected
        schema = schema_shape(expected)
    elif case_id.startswith("STOCK-"):
        required = ["report.md", "result.json"]
        data = {
            "symbol": "R001",
            "market": "FIXTURE",
            "timezone": "Asia/Shanghai",
            "currency": "CNY",
            "as_of": "2026-10-22T16:00:00+08:00",
            "receipt": nonce,
        }
        if case_id == "STOCK-01":
            data["sources"] = [
                {
                    "id": "filing",
                    "type": "official",
                    "published_at": "2026-10-22T15:00:00+08:00",
                    "fact": "产品已发布",
                },
                {
                    "id": "comment",
                    "type": "commentary",
                    "published_at": "2026-10-22T14:00:00+08:00",
                    "fact": "分析者认为需求可能增长",
                },
                {
                    "id": "future",
                    "type": "official",
                    "published_at": "2026-10-22T18:00:00+08:00",
                    "fact": "未来公告，不可引用",
                },
            ]
            expected = {
                **{
                    k: data[k]
                    for k in ("symbol", "market", "timezone", "currency", "as_of", "receipt")
                },
                "facts": ["filing"],
                "commentary": ["comment"],
                "real_time": False,
            }
            schema = {
                "facts": ["source.id"],
                "commentary": ["source.id"],
                "contract": "facts和commentary按类型分别列出截至as_of可用来源的id，不使用来源正文替代id；报告再陈述事实与评论并引用来源。",
            }
        elif case_id == "STOCK-03":
            data.update(
                product_announced=True, business_share=0.2, orders=None, profit_increment=None
            )
            expected = {
                "receipt": nonce,
                "facts": {"product_announced": True, "business_share": 0.2},
                "orders": None,
                "profit_increment": None,
                "target_price": None,
                "guaranteed_return": False,
            }
            schema = {
                "scenarios": [
                    {"condition": "非空条件，可附情景名称、影响方向与依据；至少两个情景"}
                ],
                "falsification": ["非空可证伪条件，至少一个"],
            }
            judge["text_required"] = ["情景", "反证"]
        elif case_id == "STOCK-04":
            data.update(
                status=503,
                snapshot={
                    "price": 18.8,
                    "as_of": "2026-10-21T16:00:00+08:00",
                    "currency": "CNY",
                    "market": "FIXTURE",
                },
            )
            expected = {
                "receipt": nonce,
                "status": 503,
                "price": 18.8,
                "as_of": data["snapshot"]["as_of"],
                "currency": "CNY",
                "market": "FIXTURE",
                "real_time": False,
                "order_created": False,
            }
        else:
            data.update(
                preference={"turnover": "low", "horizon": "long"},
                exclude=["X"],
                companies=[
                    {"symbol": "R001", "industry": "X"},
                    {
                        "symbol": "R002",
                        "industry": "Y",
                        "fact": "已发布年度报告",
                        "condition": "核对下季度现金流",
                    },
                    {
                        "symbol": "R003",
                        "industry": "Z",
                        "fact": "已披露研发投入",
                        "condition": "核对产品量产进展",
                    },
                ],
            )
            expected = {
                "receipt": nonce,
                "watchlist": ["R002", "R003"],
                "horizon": "long",
                "orders": [],
                "guaranteed_return": False,
            }
        files["market.json"] = json.dumps(data, ensure_ascii=False)
        judge["expected"] = expected
        schema = {**schema_shape(expected), **schema}
    elif case_id.startswith("QUANT-"):
        required = ["analysis.py", "result.json", "report.md"]
        data = {
            "bars": prices(180 if case_id == "QUANT-02" else 60),
            "initial_cash": 10000,
            "fee": 0.001,
            "slippage": 0.001,
            "receipt": nonce,
        }
        if case_id == "QUANT-01":
            expected = {**backtest(data["bars"]), "receipt": nonce}
            required += ["trades.csv", "equity.svg"]
            schema = {
                "contract": "MA3/MA5：已有5个收盘价后，MA3>MA5持有一倍现货多头，否则空仓；下一bar开盘成交，全现金含买入费买入或全部卖出；不强制期末平仓；小数股可用；result.json含trades（signal_bar/execution_bar/side/price/quantity/fee）、逐bar equity数值数组与原始receipt字符串（不得改成对象），只包含这三个顶层字段；买入价=open*(1+slippage)，卖出价=open*(1-slippage)，费用和滑点均取prices.json。另存trades.csv、equity.svg。误差1e-5。analysis.py重跑产生同结果。"
            }
        elif case_id == "QUANT-02":
            from roami_agent.improvement.evaluation.case_catalog import case_resource

            data = json.loads(case_resource(case_id, "inputs/prices.json").read_text())
            data["receipt"] = nonce
            files["tools.json"] = case_resource(case_id, "tools.json").read_text()
            validation_start, validation_end = data["validation"]
            test_start, test_end = data["test"]
            costs = {"fee": data["fee"], "slip": data["slippage"]}
            scores = []
            for fast, slow in data["candidates"]:
                scores.append(
                    (
                        backtest(
                            data["bars"][validation_start:validation_end], fast, slow, **costs
                        )["equity"][-1],
                        fast,
                        slow,
                    )
                )
            best = min(scores, key=lambda x: (-x[0], x[1], x[2]))
            expected = {
                "receipt": nonce,
                "selected": [best[1], best[2]],
                "tried": data["candidates"],
                "selection_range": data["validation"],
                "test_range": data["test"],
                "test_used_for_selection": False,
                "test_final_equity": backtest(
                    data["bars"][test_start:test_end], best[1], best[2], **costs
                )["equity"][-1],
            }
            schema = {
                "receipt": "string",
                "selected": ["fast整数", "slow整数"],
                "tried": [["fast整数", "slow整数"]],
                "selection_range": ["start整数", "end整数"],
                "test_range": ["start整数", "end整数"],
                "test_used_for_selection": "boolean",
                "test_final_equity": "number",
                "contract": "训练/验证/测试分别独立重置10000现金和指标历史，不借用前一区间价格；区间为左闭右开。已有slow个收盘价后，MAfast>MAslow持有一倍现货多头，否则空仓；收盘信号下一bar开盘执行，全现金含买入费买入或全部卖出，小数股可用，不强制期末平仓，末净值按最后收盘价估值。买入价=open*(1+slippage)，卖出价=open*(1-slippage)，费用与滑点取prices.json。用验证区间末净值选择，平局按fast/slow升序；测试仅选定后一次。receipt原样保留为字符串；selected、tried与区间使用数组，不改成对象。selection-log.json中candidates为试过的[fast,slow]数组列表、selection_range为验证区间[start,end]、test_used_for_selection=false，可另附训练和验证分数；不得计算或记录未选中候选的测试分数。误差1e-5。",
            }
            required.append("selection-log.json")
        elif case_id == "QUANT-03":
            files["flawed.py"] = (
                "signal = close.shift(-1) > close\nuniverse = [s for s in stocks if s.alive_at_end]\n"
            )
            data["dead_security"] = {"symbol": "DELISTED", "return": -0.8}
            expected = {
                "receipt": nonce,
                "biases": ["lookahead", "survivorship"],
                "old_claim_valid": False,
                "delisted_included": True,
            }
            strategy = backtest(data["bars"])
            judge["corrected_expected"] = {
                "receipt": nonce,
                "universe": ["REGULAR", "DELISTED"],
                "strategy": strategy,
                "delisted_final_equity": 2000,
                "portfolio_final_equity": (strategy["equity"][-1] + 2000) / 2,
            }
            schema = {
                "contract": "result.json含receipt、biases=[lookahead,survivorship]、old_claim_valid=false、delisted_included=true；报告定位flawed.py两行。analysis.py重跑生成corrected.json：receipt、universe=[REGULAR,DELISTED]、strategy（REGULAR的因果MA3/MA5回测，已有5个收盘价后MA3>MA5持有一倍现货多头，否则空仓；收盘信号下一bar开盘全现金含买入费买入或全部卖出，小数股可用，初始10000，费用与滑点取prices.json，买入价=open*(1+slippage)、卖出价=open*(1-slippage)；含trades和逐bar equity数值数组，不强制期末平仓）、delisted_final_equity=2000、portfolio_final_equity（两资产初始各10000、等权归一至10000，取两者末净值平均）。trades每项字段为signal_bar/execution_bar/side/price/quantity/fee；receipt原样保留字符串。不能只写已修复。"
            }
            schema["contract"] += (
                "strategy的必要字段为trades、equity；可附description及initial_cash、fee、slippage、final_cash、final_quantity、final_equity摘要，数值必须与输入和逐笔交易、净值一致，不接受其他额外字段。"
            )
            required.append("corrected.json")
        elif case_id == "QUANT-04":
            data = {
                "equity": [100, 110, 99, 120],
                "benchmark": [100, 105, 103, 115],
                "trades": [{"notional": 100, "side": "buy"}, {"notional": 120, "side": "sell"}],
                "fees": [0.001, 0.003],
                "receipt": nonce,
            }
            expected = {
                "receipt": nonce,
                "zero_cost_return": 0.2,
                "max_drawdown": 0.1,
                "benchmark_return": 0.15,
                "costs": [0.22, 0.66],
                "net_returns": [0.1978, 0.1934],
            }
            schema = {
                "contract": "逐笔费用=notional*fee，净收益=(120-总费用)/100-1；result.json按给定字段计算：receipt,zero_cost_return,max_drawdown,benchmark_return,costs,net_returns；所有数值用比例，误差1e-5。"
            }
        else:
            data.update(
                missing_bars=[8, 13],
                halted_bars=[14],
                split={"bar": 30, "ratio": 2},
                execution_prices="raw",
                signal_prices="adjusted",
                unresolved=True,
            )
            expected = {
                "receipt": nonce,
                "completed": False,
                "unresolved": True,
                "trades_on_unavailable_bars": [],
                "execution_prices": "raw",
                "signal_prices": "adjusted",
                "missing_bars": [8, 13],
                "halted_bars": [14],
                "split_ratio": 2,
            }
            schema = {
                **schema_shape(expected),
                "contract": "result.json原样复制prices.json中的字符串字段receipt、execution_prices、signal_prices；不生成新的哈希凭证或在这些字段内附加说明，解释放入report.md。",
            }
        files["prices.json"] = json.dumps(data)
        if case_id in {"QUANT-01", "QUANT-03"}:
            schema["contract"] += (
                " trades 中 side 必须是小写字符串 buy 或 sell，不使用大写或其他别名。"
            )
        judge.update(expected=expected, numeric_tolerance=1e-5)
    elif case_id == "LAB-05":
        files["experience-v1.json"] = json.dumps(
            {"version": 1, "entry": "python test.py", "source": "project/README.md"}
        )
        files["project/README.md"] = (
            "测试入口是 python verify_current.py，旧的 python test.py 已停用。版本凭证 " + nonce
        )
        files["project/verify_current.py"] = "print('VERIFIED:" + nonce + "')\n"
        files["project/test.py"] = "raise SystemExit('RETIRED_ENTRY')\n"
        required = ["experience.json", "transfer-log.json", "result.json"]
        expected = {
            "receipt": nonce,
            "entry": "python verify_current.py",
            "old_entry_used": False,
            "source": "project/README.md",
            "revalidated": True,
        }
        schema = {
            **schema_shape(expected),
            "contract": "experience-v1.json是已过期经验；当前project为新版本。核对新README并运行新入口，再保存experience.json。receipt为README版本凭证的原始字符串；source为project/README.md路径，不附加说明。transfer-log.json用output或stdout字符串保存新入口的原始标准输出；若两字段同时存在，必须一致。这是已有经验失效后的复核任务，不是两次独立会话的长期记忆测量。不得把本次题目答案作为经验。",
        }
        judge["expected"] = expected
    else:
        raise ValueError(case_id)
    instructions = schema.pop("contract", None)
    prompt = "\n执行合同：所有输入是离线合成夹具，不是实时新闻或行情。" + json.dumps(
        schema, ensure_ascii=False
    )
    if instructions:
        prompt += "\n执行要求（不是结果字段）：" + instructions
    prompt += (
        "\n必要产物："
        + ", ".join(required)
        + "。如有result.json，字段按合同给出，不添加未证实信息。报告引用输入文件并原样包含输入的receipt字符串；有source_receipt时同样原样保留。"
    )
    editable = {"app.py", "index.html", "existing.js"}
    protected = sorted(set(protected) | (set(files) - editable))
    judge.update(
        required=required,
        protected={k: digest(files[k]) for k in protected},
        protected_assets={
            k: hashlib.sha256(base64.b64decode(v)).hexdigest() for k, v in assets.items()
        },
    )
    return {
        "files": files,
        "asset_files": assets,
        "judge": judge,
        "prompt_suffix": prompt,
        "fixture_receipt": nonce,
        "fixture_revision": "roami-authored-v3",
        "contract_revision": (
            "authored-v15-input-field-serialization"
            if case_id == "QUANT-05"
            else "authored-v14-explicit-trade-sides"
            if case_id in {"QUANT-01", "QUANT-03"}
            else "authored-v13-visible-quant-inputs"
            if case_id == "QUANT-02"
            else "authored-v11-verified-backtest-summaries"
            if case_id == "QUANT-03"
            else "authored-v10-separated-instructions"
        )
        if instructions
        else None,
    }


def schema_shape(expected):
    if isinstance(expected, dict):
        return {k: schema_shape(v) for k, v in expected.items()}
    if isinstance(expected, list):
        return [schema_shape(expected[0])] if expected else []
    return (
        "string"
        if isinstance(expected, str)
        else "boolean"
        if type(expected) is bool
        else "number"
        if isinstance(expected, (int, float))
        else "null_or_unavailable"
    )


CSV_SERVER = """from http.server import BaseHTTPRequestHandler,HTTPServer
import os,json
seen={};count=0;attempts=0
class H(BaseHTTPRequestHandler):
 def do_GET(self):
  if self.path=="/stats": body=json.dumps({"rows":count,"attempts":attempts}).encode();self.send_response(200);self.end_headers();self.wfile.write(body);return
  self.send_response(200);self.send_header("Content-Type","text/html");self.end_headers();self.wfile.write(open("index.html","rb").read())
 def do_POST(self):
  global count,attempts
  attempts+=1;data=json.loads(self.rfile.read(int(self.headers["Content-Length"])));key=self.headers.get("Idempotency-Key","")
  if not key: self.send_response(400);self.end_headers();return
  if key not in seen: count+=len(data["rows"]);seen[key]=data
  self.send_response(503 if attempts==2 else 200);self.end_headers();self.wfile.write(json.dumps({"rows":count}).encode())
if __name__=="__main__":HTTPServer(("127.0.0.1",int(os.getenv("PORT","8000"))),H).serve_forever()
"""
