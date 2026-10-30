"""An optional package uses existing Skill, shell and delivery contracts unchanged."""

import base64
import copy
import hashlib
import io
import json
import sys
from contextlib import asynccontextmanager
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches
from test_runtime import until

from roami_agent.app import Runtime
from roami_agent.protocol.errors import DomainError
from roami_agent.protocol.models import ModelEvent
from roami_agent.storage.database import row


async def test_external_skill_import_executes_and_delivers_with_unchanged_foundations(
    store, tmp_path, monkeypatch
):
    package = Path(__file__).resolve().parents[3] / "skills/presentation-edit"
    workspace = tmp_path / "work with spaces"
    workspace.mkdir()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for text, y in [("Body", 4), ("Heading", 1)]:
        shape = slide.shapes.add_textbox(Inches(1), Inches(y), Inches(4), Inches(1))
        shape.text = text
        shape.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    source = workspace / "source.pptx"
    deck.save(source)
    original = source.read_bytes()
    calls = []

    class Main:
        model = "skill-package-main-contract"

        async def stream(self, messages, tools):
            yield ModelEvent("text", "File saved.")

    class Worker:
        model = "skill-package-worker-contract"

        async def stream(self, messages, tools):
            receipts = [json.loads(m["content"]) for m in messages if m["role"] == "tool"]
            if not receipts:
                name, args = "read", {"path": "skill://presentation-edit"}
            elif len(receipts) < 4:
                script = str(Path(receipts[0]["packagePath"]) / "scripts/presentation.py")
                if len(receipts) == 2:
                    observed = json.loads(receipts[-1]["stdout"])
                    request = {
                        "path": "source.pptx",
                        "expectedSha256": observed["sha256"],
                        "destination": "edited.pptx",
                        "expectedDestinationSha256": None,
                        "operations": [
                            {
                                "slide": 1,
                                "target": {"order": "top_to_bottom", "ordinal": 1},
                                "alignment": "center",
                            }
                        ],
                    }
                    argv = ["edit", "--request-json", json.dumps(request)]
                else:
                    argv = [
                        "inspect",
                        "--path",
                        "source.pptx" if len(receipts) == 1 else "edited.pptx",
                    ]
                name, args = (
                    "bash",
                    {"argv": [sys.executable, script, *argv, "--workspace", str(workspace)]},
                )
            else:
                checked = json.loads(receipts[-1]["stdout"])
                assert (
                    checked["slides"][0]["textObjects"][1]["paragraphs"][0]["alignment"] == "center"
                )
                yield ModelEvent("text", "已验证：[edited.pptx](edited.pptx)")
                return
            calls.append(name)
            yield ModelEvent(
                "tool",
                {
                    "id": str(len(receipts)),
                    "function": {"name": name, "arguments": json.dumps(args)},
                },
            )

    @asynccontextmanager
    async def prepare(*args, **kwargs):
        yield {"workspaceId": "external-skill", "directory": str(workspace)}

    runtime = Runtime(store, Main(), worker_model=Worker(), memory_model=Main())
    runtime.dispatch.native_executor.workspace_factory = prepare
    before = copy.deepcopy(runtime.tools.catalog("worker"))
    tool_names = set(runtime.tools.registry)
    implicit_reads = []
    original_read = runtime.files.read

    async def observe_read(*args, **kwargs):
        if not calls:
            implicit_reads.append(args[0])
        return await original_read(*args, **kwargs)

    monkeypatch.setattr(runtime.files, "read", observe_read)
    assert "presentation-edit" not in runtime.tools.skills
    payload = {
        "name": "presentation-edit",
        "description": next(
            line.partition(":")[2].strip()
            for line in (package / "SKILL.md").read_text().splitlines()
            if line.startswith("description:")
        ),
        "text": (package / "SKILL.md").read_text(),
        "files": {
            str(path.relative_to(package)): base64.b64encode(path.read_bytes()).decode()
            for folder in ["scripts", "references"]
            for path in (package / folder).rglob("*")
            if path.is_file() and "__pycache__" not in path.parts
        },
    }
    await runtime.capability_imports.save("local", "skills", payload)
    assert runtime.tools.catalog("worker") == before and set(runtime.tools.registry) == tool_names
    assert any(
        item["name"] == "presentation-edit"
        for item in runtime.tools.directory.list("幻灯片", "skill")["items"]
    )
    await runtime.start()
    try:
        conversation = await runtime.conversations.create("alice")
        task = await runtime.tasks.create(
            "alice",
            conversation["id"],
            "Use the external Skill to inspect source.pptx and center its top heading in edited.pptx.",
            key="external-skill",
            target={"kind": "new_session", "directory": str(workspace)},
            record_intent=True,
        )
        runtime.wakeup.set()
        await until(
            store,
            lambda db: row(
                db,
                "SELECT id FROM turns WHERE conversation_id=? AND trigger_kind='completion' AND state='completed'",
                (conversation["id"],),
            ),
        )
        assert implicit_reads == []
        assert calls == ["read", "bash", "bash", "bash"]
        references = await runtime.artifacts.list("alice", task_id=task["id"])
        assert len(references) == 1 and references[0]["name"] == "edited.pptx"
        artifact = (await runtime.artifacts.read("alice", references[0]["id"]))[1]
        result = Presentation(io.BytesIO(artifact))
        assert result.slides[0].shapes[0].text_frame.paragraphs[0].alignment == PP_ALIGN.LEFT
        assert result.slides[0].shapes[1].text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
        assert (
            hashlib.sha256(original).hexdigest() == hashlib.sha256(source.read_bytes()).hexdigest()
        )
        (workspace / "edited.pptx").unlink()
        assert (await runtime.artifacts.read("alice", references[0]["id"]))[1] == artifact
        with pytest.raises(DomainError, match="ARTIFACT_NOT_FOUND"):
            await runtime.artifacts.read("bob", references[0]["id"])
        await runtime.capability_settings.save(
            "local", {"kind": "skills", "name": "presentation-edit", "enabled": False}
        )
        with pytest.raises(DomainError, match="SKILL_DISABLED"):
            await runtime.tools.invoke(
                {
                    "function": {
                        "name": "read",
                        "arguments": json.dumps({"path": "skill://presentation-edit"}),
                    }
                },
                "worker",
                {},
            )
        await runtime.capability_imports.save(
            "local", "skills", None, name="presentation-edit", remove=True
        )
        assert "presentation-edit" not in runtime.tools.skills
        assert (
            runtime.tools.catalog("worker") == before and set(runtime.tools.registry) == tool_names
        )
    finally:
        await runtime.close()
