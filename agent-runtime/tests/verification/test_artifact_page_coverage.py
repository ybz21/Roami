"""Presentation evidence uses the existing combined image allowance without dropping pages."""

import hashlib
import os

import pytest
from pptx import Presentation
from pptx.util import Inches

from roami_agent.improvement.evaluation.artifact_evidence import collect_artifact_evidence
from roami_agent.improvement.evaluation.execution.containers import ContainerExecutor
from roami_agent.protocol.errors import DomainError
from roami_agent.server.configuration import Settings


@pytest.mark.parametrize("source_pages,updated_pages", [(1, 9), (9, 1), (8, 8), (8, 9)])
async def test_all_pages_fit_the_combined_allowance_or_no_evidence_is_published(
    tmp_path, source_pages, updated_pages
):
    image = os.environ.get("ROAMI_RSI_TEST_IMAGE")
    if not image:
        pytest.skip("Set ROAMI_RSI_TEST_IMAGE for real LibreOffice and PDF rendering")
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    for name, count in [("source.pptx", source_pages), ("updated.pptx", updated_pages)]:
        deck = Presentation()
        for index in range(count):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(7), Inches(1)
            ).text = f"Controlled page {index + 1}"
        deck.save(workspace / name)
    (workspace / "changes.md").write_text(
        "Controlled rendering fixture, not an RSI quality result."
    )
    files = [workspace / name for name in ["source.pptx", "updated.pptx", "changes.md"]]
    before = {f.name: hashlib.sha256(f.read_bytes()).hexdigest() for f in files}
    case = {"id": "DOC-03", "prompt": "Review all supplied pages.", "judge": {"kind": "authored"}}
    grade = {"checks": {"controlled_fixture": True}, "evidence": []}
    executor = ContainerExecutor(Settings(), image=image)
    if source_pages + updated_pages > 16:
        with pytest.raises(DomainError):
            await collect_artifact_evidence(case, grade, tmp_path, executor)
        assert not (tmp_path / "artifact-evidence.json").exists()
    else:
        result = await collect_artifact_evidence(case, grade, tmp_path, executor)
        labels = [f"source.pptx page {i + 1}" for i in range(source_pages)] + [
            f"updated.pptx page {i + 1}" for i in range(updated_pages)
        ]
        assert [image["label"] for image in result["images"]] == labels
        assert all(item["sha256"] and item["width"] and item["height"] for item in result["images"])
        assert {item["file"]: item["sha256"] for item in result["artifacts"]} == before
        assert (tmp_path / "artifact-evidence.json").is_file()
    assert before == {f.name: hashlib.sha256(f.read_bytes()).hexdigest() for f in files}
