"""Five plots need distinct pages; additional cover pages and editable charts are valid."""

import base64
import hashlib
import io
import json
import os
import shutil

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from roami_agent.improvement.evaluation.benchmarks.stock_tasks import (
    PRESENTATION_CHECK,
    grade_stock,
)
from roami_agent.improvement.evaluation.execution.containers import ContainerExecutor
from roami_agent.improvement.evaluation.execution.stock_environment import OUTPUTS
from roami_agent.server.configuration import Settings


def picture():
    output = io.BytesIO()
    Image.new("RGB", (10, 10), "white").save(output, format="PNG")
    output.seek(0)
    return output


def deck(path, kind, count=5, *, one_page=False):
    result = Presentation()
    result.slides.add_slide(result.slide_layouts[6])
    shared = result.slides.add_slide(result.slide_layouts[6]) if one_page else None
    for _ in range(count):
        if kind == "placeholder":
            result.slides.add_slide(result.slide_layouts[8]).placeholders[1].insert_picture(
                picture()
            )
            continue
        slide = shared if shared is not None else result.slides.add_slide(result.slide_layouts[6])
        shapes = slide.shapes
        if kind == "group":
            shapes = shapes.add_group_shape().shapes.add_group_shape().shapes
        if kind == "chart":
            data = CategoryChartData()
            data.categories = ["first", "second"]
            data.add_series("series", [1, 2])
            shapes.add_chart(XL_CHART_TYPE.LINE, Inches(1), Inches(1), Inches(4), Inches(3), data)
        elif kind != "text":
            shapes.add_picture(picture(), Inches(1), Inches(1))
        else:
            shapes.add_textbox(
                Inches(1), Inches(1), Inches(4), Inches(1)
            ).text = "Summary without a plot"
    result.slides.add_slide(result.slide_layouts[6])
    result.save(path)


def count_plots(path):
    namespace = {}
    exec(compile(PRESENTATION_CHECK, "frozen_stock_presentation.py", "exec"), namespace)  # noqa: S102 - trusted in-repository evaluator
    return namespace["plot_slide_count"](path)


@pytest.mark.parametrize("kind", ["picture", "chart", "group", "placeholder"])
def test_five_plot_pages_allow_cover_pages_and_editable_or_grouped_visuals(tmp_path, kind):
    path = tmp_path / "deck.pptx"
    deck(path, kind)
    assert len(Presentation(path).slides) == 7
    assert count_plots(path) == 5


@pytest.mark.parametrize(
    ("kind", "count", "one_page", "expected"),
    [("picture", 4, False, 4), ("picture", 5, True, 1), ("text", 5, False, 0)],
)
def test_missing_plot_pages_cannot_pass_by_adding_text_pages_or_many_images_on_one_page(
    tmp_path, kind, count, one_page, expected
):
    path = tmp_path / "deck.pptx"
    deck(path, kind, count, one_page=one_page)
    assert count_plots(path) == expected < 5


@pytest.mark.parametrize("correct_csv", [True, False])
async def test_real_stock_program_keeps_csv_checks_and_independent_quality(tmp_path, correct_csv):
    image = os.environ.get("ROAMI_RSI_TEST_IMAGE")
    if not image or not shutil.which("docker"):
        pytest.skip("Requires the frozen RSI image with pandas and python-pptx")
    remote = tmp_path / "stock-remote"
    remote.mkdir()
    expected = b"Date,Close\n2026-01-01,10\n2026-01-02,11\n"
    (remote / "stock_cleaned.csv").write_bytes(
        expected if correct_csv else expected.replace(b",11", b",99")
    )
    deck(remote / "stock_analysis.pptx", "chart")
    (remote / "stock_analysis_script.md").write_text(
        "Presenter script for the five requested plots. " * 4
    )
    state = {
        "initial": {"source_sha256": "source", "outputs": {}},
        "final": {
            "source_sha256": "source",
            "outputs": {
                name: {"sha256": hashlib.sha256((remote / name).read_bytes()).hexdigest()}
                for name in OUTPUTS
            },
        },
    }
    (tmp_path / "stock-snapshot.json").write_text(json.dumps(state))
    executor = ContainerExecutor(Settings(), image=image)

    class Readonly:
        async def sandbox(self, workspace, program, **kwargs):
            return await executor.sandbox(workspace, program, copy_workspace=True, **kwargs)

    case = {
        "judge": {
            "raw_sha256": "source",
            "golden": base64.b64encode(expected).decode(),
            "checkpoint": "def grade_checkpoint1():\n return pd.read_csv('/workspace/stock_cleaned.csv').equals(pd.read_csv('/utils/stock_data.csv'))",
        }
    }
    grade = await grade_stock(case, Readonly(), root=tmp_path)
    assert grade["checks"]["five_plot_slides"]
    assert grade["checks"]["adapted_csv_values"] is correct_csv
    assert grade["hard_passed"] is correct_csv and grade["pending"] is correct_csv
    assert grade["native_partial"]["scored_points"] == int(correct_csv)
    assert grade["rubric_required"] and not grade["passed"] and grade["quality"] == 0
