"""Ordinary native exports need a writable XDG config inside the disposable filesystem."""

import hashlib
import json
import os

import pytest
from pptx import Presentation

from roami_agent.improvement.domain.contracts import digest
from roami_agent.improvement.evaluation.execution import containers
from roami_agent.improvement.evaluation.protocol_suite import protocol_suite
from roami_agent.server.configuration import Settings


@pytest.mark.parametrize("runtime", [False, True])
async def test_ordinary_office_export_works_without_a_private_profile_argument(tmp_path, runtime):
    image = os.environ.get("ROAMI_RSI_TEST_IMAGE")
    if not image:
        pytest.skip("Set ROAMI_RSI_TEST_IMAGE for real LibreOffice export")
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    source = workspace / "source.pptx"
    deck.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    program = """import hashlib,json,pathlib,subprocess
pathlib.Path.home().mkdir(parents=True,exist_ok=True)
output=pathlib.Path('/tmp/exports');output.mkdir()
result=subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',str(output),'/workspace/source.pptx'],capture_output=True,timeout=60)
pdf=output/'source.pdf'
print(json.dumps({'exit':result.returncode,'stderr':result.stderr.decode(),'pdf':pdf.is_file() and pdf.read_bytes().startswith(b'%PDF-'),'source':hashlib.sha256(pathlib.Path('/workspace/source.pptx').read_bytes()).hexdigest()}))
"""
    executor = containers.ContainerExecutor(Settings(), image=image)
    result = json.loads(await executor.sandbox(workspace, program, runtime=runtime, timeout=75))
    assert result["exit"] == 0, result
    assert result["pdf"] is True
    assert result["source"] == hashlib.sha256(source.read_bytes()).hexdigest() == before


async def test_actor_launch_uses_the_disposable_config_directory(tmp_path, monkeypatch):
    created = []

    async def docker(*args, **kwargs):
        if args[0] == "create":
            created.append(args)
        if args[0] == "start":
            raise TimeoutError
        return b""

    async def remove(*args):
        pass

    monkeypatch.setattr(containers, "docker_command", docker)
    executor = containers.ContainerExecutor(Settings(), shared_daily_limit=True)
    executor.image_id = "no-model-fixture"
    monkeypatch.setattr(executor, "remove", remove)
    await executor.run(protocol_suite()["cases"][0], {}, tmp_path)
    assert len(created) == 1
    args = created[0]
    environment = dict(arg.split("=", 1) for i, arg in enumerate(args) if i and args[i - 1] == "-e")
    assert environment["XDG_CONFIG_HOME"] == "/tmp"
    assert environment["HOME"].startswith("/tmp/")
    assert "--read-only" in args and args[args.index("--network") + 1] == "none"


async def test_office_font_uses_the_metric_compatible_substitute(tmp_path):
    image = os.environ.get("ROAMI_RSI_TEST_IMAGE")
    if not image:
        pytest.skip("Set ROAMI_RSI_TEST_IMAGE for real Office font resolution")
    executor = containers.ContainerExecutor(Settings(), image=image)
    fingerprint = await executor.preflight()
    assert fingerprint["environment_hash"] != digest(
        [
            fingerprint["image_id"],
            fingerprint["artifact_renderer_image_id"],
            "network-none",
            "nonroot-readonly-v1",
        ]
    )
    program = "import subprocess;print(subprocess.check_output(['fc-match','-f','%{family}','Calibri'],text=True))"
    family = (await executor.sandbox(tmp_path, program)).decode().strip()
    assert family == "Carlito", family
