"""Focused inspection keeps original identity and returns only requested objects."""

import hashlib
import json
from pathlib import Path

import pytest
from document_publication import Workspace
from pptx import Presentation
from pptx.util import Inches
from pptx_package import PresentationError
from presentation import inspect, main


def deck(root):
    presentation = Presentation()
    case = Path(__file__).resolve().parents[3] / "cases/case-64-mech-dp01"
    recipe = json.loads((case / "input/focused-slides.json").read_text())
    for entry in recipe["slides"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
        box.text = entry["label"] + entry["text"] * entry["repeat"]
    path = root / "deck.pptx"
    presentation.save(path)
    return path


def test_single_slide_keeps_deck_hash_and_complete_object_metadata(tmp_path):
    path = deck(tmp_path)
    before = path.read_bytes()
    full = inspect(Workspace(tmp_path), path.name)
    selected = inspect(Workspace(tmp_path), path.name, 3)
    assert selected["slides"] == [full["slides"][2]]
    assert selected["slideCount"] == 12
    assert selected["sha256"] == hashlib.sha256(before).hexdigest()
    assert selected["editorPalette"] == full["editorPalette"]
    assert len(json.dumps(selected)) < len(json.dumps(full)) / 4
    assert path.read_bytes() == before


@pytest.mark.parametrize("slide", [0, -1, 13, True])
def test_unknown_or_noninteger_slide_does_not_mutate(tmp_path, slide):
    path = deck(tmp_path)
    before = path.read_bytes()
    with pytest.raises(PresentationError, match="PRESENTATION_TARGET_NOT_FOUND"):
        inspect(Workspace(tmp_path), path.name, slide)
    assert path.read_bytes() == before


def test_cli_focus_returns_the_requested_slide(tmp_path, capsys):
    path = deck(tmp_path)
    assert (
        main(
            [
                "inspect",
                "--workspace",
                str(tmp_path),
                "--path",
                path.name,
                "--slide",
                "2",
            ]
        )
        == 0
    )
    result = json.loads(capsys.readouterr().out)
    assert result["ok"] and result["slideCount"] == 12
    assert [slide["number"] for slide in result["slides"]] == [2]
