import io

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches


def deck_bytes(*, same_position=False):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    body = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(5), Inches(1))
    body.text = "Body"
    heading = slide.shapes.add_textbox(
        Inches(1), Inches(4 if same_position else 1), Inches(5), Inches(1)
    )
    heading.text = "Heading"
    for shape in [body, heading]:
        shape.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        shape.text_frame.paragraphs[0].runs[0].font.size = Inches(0.3)
    stream = io.BytesIO()
    deck.save(stream)
    return stream.getvalue()
