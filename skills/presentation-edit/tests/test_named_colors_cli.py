"""Named palette edits do not depend on a model guessing RGB values."""

import hashlib
import json

import office_palette
import presentation
import pytest
from deck_fixture import deck_bytes
from pptx import Presentation


@pytest.fixture
def color_case(tmp_path, monkeypatch):
    source = tmp_path / "input.pptx"
    source.write_bytes(deck_bytes())
    palette = tmp_path / "standard.soc"
    palette.write_text(
        '<colors xmlns:draw="urn:palette"><color draw:name="Blue" draw:color="#123456"/><color draw:name="Orange" draw:color="#AB7654"/></colors>'
    )
    monkeypatch.setattr(office_palette, "PALETTES", (palette,))
    return source, palette


def arguments(source, names, sha=None):
    return [
        "colors",
        "--workspace",
        str(source.parent),
        "--path",
        source.name,
        "--destination",
        "updated.pptx",
        "--expected-sha256",
        sha or hashlib.sha256(source.read_bytes()).hexdigest(),
        "--slide",
        "1",
        "--names",
        *names,
    ]


def test_named_colors_follow_geometry_and_actual_installed_palette(color_case, capsys):
    source, _ = color_case
    original = source.read_bytes()
    assert presentation.main(arguments(source, ["Blue", "Orange"])) == 0
    receipt = json.loads(capsys.readouterr().out)
    assert receipt["ok"] and receipt["writePublished"]
    deck = Presentation(source.with_name("updated.pptx"))
    assert (
        str(deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.color.rgb)
        == "AB7654"
    )
    assert (
        str(deck.slides[0].shapes[1].text_frame.paragraphs[0].runs[0].font.color.rgb)
        == "123456"
    )
    assert source.read_bytes() == original


@pytest.mark.parametrize("failure", ["unknown_name", "count", "source", "palette"])
def test_invalid_named_color_batch_publishes_nothing(color_case, capsys, failure):
    source, palette = color_case
    original = source.read_bytes()
    names, sha = ["Blue", "Orange"], None
    if failure == "unknown_name":
        names[1] = "Unknown"
    elif failure == "count":
        names.pop()
    elif failure == "source":
        sha = "0" * 64
    else:
        palette.unlink()
    assert presentation.main(arguments(source, names, sha)) == 1
    receipt = json.loads(capsys.readouterr().out)
    assert receipt["ok"] is False and receipt["writePublished"] is False
    assert not source.with_name("updated.pptx").exists()
    assert source.read_bytes() == original
