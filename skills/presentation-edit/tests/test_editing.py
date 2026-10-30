"""Standalone package execution and document publication, without Agent runtime imports."""

import hashlib
import io
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path

import office_palette
import presentation
import pytest
from deck_fixture import deck_bytes
from document_publication import PublicationError, Workspace
from pptx import Presentation
from pptx.util import Inches
from pptx_package import PresentationError


@pytest.fixture
def document(tmp_path, monkeypatch):
    source = tmp_path / "source.pptx"
    source.write_bytes(deck_bytes())
    palette = tmp_path / "standard.soc"
    palette.write_text(
        '<colors xmlns:draw="urn:palette"><color draw:name="Example" draw:color="#12ab34"/></colors>'
    )
    monkeypatch.setattr(office_palette, "PALETTES", (palette,))
    workspace = Workspace(tmp_path)
    observed = presentation.inspect(workspace, "source.pptx")
    request = {
        "path": "source.pptx",
        "expectedSha256": observed["sha256"],
        "destination": "edited.pptx",
        "expectedDestinationSha256": None,
        "operations": [
            {
                "slide": 1,
                "target": {"order": "top_to_bottom", "ordinal": 1},
                "alignment": "right",
            }
        ],
    }
    return workspace, source, request


def test_standalone_edit_preserves_source_other_parts_and_font(document):
    workspace, source, request = document
    original = source.read_bytes()
    request["operations"][0]["color"] = {
        "editor": "LibreOffice Impress",
        "palette": "Standard",
        "name": "Example",
    }
    result = presentation.edit(workspace, request)
    output = source.with_name("edited.pptx")
    assert result["writePublished"] is True
    assert result["sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert source.read_bytes() == original
    body, heading = presentation.inspect(workspace, "edited.pptx")["slides"][0][
        "textObjects"
    ]
    assert body["paragraphs"][0]["alignment"] == "left"
    assert heading["paragraphs"][0]["alignment"] == "right"
    assert heading["paragraphs"][0]["runs"][0]["color"]["value"] == "12AB34"
    with (
        zipfile.ZipFile(io.BytesIO(original)) as before,
        zipfile.ZipFile(output) as after,
    ):
        assert before.namelist() == after.namelist()
        assert [
            name for name in before.namelist() if before.read(name) != after.read(name)
        ] == ["ppt/slides/slide1.xml"]
    assert Presentation(output).slides[0].shapes[1].text_frame.paragraphs[0].runs[
        0
    ].font.size == Inches(0.3)


@pytest.mark.parametrize("changed", ["source", "destination"])
def test_cas_rejects_changed_files(document, changed):
    workspace, source, request = document
    path = source if changed == "source" else source.with_name("edited.pptx")
    path.write_bytes(b"user changed this")
    with pytest.raises(PublicationError, match="CHANGED"):
        presentation.edit(workspace, request)
    assert path.read_bytes() == b"user changed this"


def test_observed_destination_can_be_revised(document):
    workspace, source, request = document
    original = source.read_bytes()
    output = source.with_name("edited.pptx")
    output.write_bytes(b"previous output")
    request["expectedDestinationSha256"] = hashlib.sha256(
        output.read_bytes()
    ).hexdigest()
    result = presentation.edit(workspace, request)
    assert result["previousSha256"] == request["expectedDestinationSha256"]
    assert source.read_bytes() == original


def test_source_change_during_fsync_prevents_publication(document, monkeypatch):
    workspace, source, request = document
    original_fsync = os.fsync
    changed = False

    def concurrent_edit(descriptor):
        nonlocal changed
        if not changed:
            changed = True
            source.write_bytes(b"new user input")
        return original_fsync(descriptor)

    monkeypatch.setattr(os, "fsync", concurrent_edit)
    with pytest.raises(PublicationError, match="SOURCE_CHANGED"):
        presentation.edit(workspace, request)
    assert source.read_bytes() == b"new user input"
    assert not source.with_name("edited.pptx").exists()
    assert not list(source.parent.glob(".presentation-*"))


def test_invalid_later_operation_never_publishes_partial_result(document):
    workspace, source, request = document
    request["operations"].append(
        {"slide": 99, "target": {"shapePath": [0]}, "alignment": "center"}
    )
    with pytest.raises(PresentationError):
        presentation.edit(workspace, request)
    assert not source.with_name("edited.pptx").exists()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == request["expectedSha256"]


@pytest.mark.parametrize(
    "kind", ["source_link", "destination_link", "parent_link", "in_place", "outside"]
)
def test_workspace_and_source_boundaries(document, kind):
    workspace, source, request = document
    original = source.read_bytes()
    outside = source.parent.parent / (source.parent.name + "-private.pptx")
    outside.write_bytes(b"private")
    if kind == "source_link":
        source.with_name("linked.pptx").symlink_to(source)
        request["path"] = "linked.pptx"
    elif kind == "destination_link":
        source.with_name("edited.pptx").symlink_to(outside)
    elif kind == "parent_link":
        (source.parent / "linked").symlink_to(outside.parent, target_is_directory=True)
        request["destination"] = "linked/edited.pptx"
    elif kind == "in_place":
        request["destination"] = request["path"]
    else:
        request["destination"] = str(outside)
    with pytest.raises((OSError, PublicationError)):
        presentation.edit(workspace, request)
    assert source.read_bytes() == original and outside.read_bytes() == b"private"


def test_tied_positions_require_an_explicit_object(document):
    workspace, source, request = document
    source.write_bytes(deck_bytes(same_position=True))
    request["expectedSha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    with pytest.raises(PresentationError, match="AMBIGUOUS"):
        presentation.edit(workspace, request)
    request["operations"][0]["target"] = {"shapePath": [1]}
    presentation.edit(workspace, request)
    assert (
        presentation.inspect(workspace, "edited.pptx")["slides"][0]["textObjects"][1][
            "paragraphs"
        ][0]["alignment"]
        == "right"
    )


@pytest.mark.parametrize(
    "spec",
    [
        {"editor": "Unknown", "palette": "Standard", "name": "Example"},
        {"editor": "LibreOffice Impress", "palette": "Standard", "name": "Missing"},
    ],
)
def test_named_colors_never_guess_fallbacks(document, spec):
    workspace, source, request = document
    request["operations"][0]["color"] = spec
    with pytest.raises(PresentationError):
        presentation.edit(workspace, request)
    assert not source.with_name("edited.pptx").exists()


def test_explicit_rgb_does_not_require_an_editor_palette(document, monkeypatch):
    workspace, _, request = document
    monkeypatch.setattr(office_palette, "PALETTES", ())
    request["operations"][0]["color"] = {"rgb": "#abcdef"}
    assert presentation.edit(workspace, request)["changes"][0]["color"] == {
        "rgb": "#ABCDEF"
    }


def test_repeated_property_never_silently_overrides(document):
    workspace, source, request = document
    request["operations"].append(
        {"slide": 1, "target": {"shapePath": [1]}, "alignment": "left"}
    )
    with pytest.raises(PresentationError, match="CONFLICT"):
        presentation.edit(workspace, request)
    assert not source.with_name("edited.pptx").exists()


def test_grouped_object_uses_observed_path(document):
    workspace, source, request = document
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    for text, y in [("Body", 4), ("Heading", 1)]:
        group.shapes.add_textbox(Inches(1), Inches(y), Inches(4), Inches(1)).text = text
    group.rotation = 30
    deck.save(source)
    observed = presentation.inspect(workspace, "source.pptx")
    request["expectedSha256"] = observed["sha256"]
    result = presentation.edit(workspace, request)
    assert (
        result["changes"][0]["shapePath"]
        == observed["slides"][0]["topToBottomPaths"][0]
        == [0, 1]
    )
    assert (
        Presentation(source.with_name("edited.pptx")).slides[0].shapes[0].rotation == 30
    )


def test_real_cli_inspects_edits_and_reopens_a_complete_package(document):
    workspace, source, request = document
    (source.parent / "changes.json").write_text(json.dumps(request))
    script = Path(presentation.__file__)
    for argv in [
        ["edit", "--request", "changes.json"],
        ["inspect", "--path", "edited.pptx"],
    ]:
        result = subprocess.run(
            [sys.executable, str(script), *argv, "--workspace", str(workspace.root)],
            text=True,
            capture_output=True,
            check=False,
            env={"PATH": os.environ["PATH"], "PYTHONNOUSERSITE": "1"},
        )
        assert result.returncode == 0, result.stderr
        assert json.loads(result.stdout)["ok"] is True
    assert (
        Presentation(source.with_name("edited.pptx")).slides[0].shapes[1].text
        == "Heading"
    )


@pytest.mark.parametrize("body", [b"not a zip", b"", b"PK\x03\x04broken"])
def test_cli_invalid_document_is_a_known_failure(document, capsys, body):
    workspace, source, _ = document
    source.write_bytes(body)
    assert (
        presentation.main(
            ["inspect", "--workspace", str(workspace.root), "--path", source.name]
        )
        == 1
    )
    result = json.loads(capsys.readouterr().out)
    assert result["ok"] is False and result["writePublished"] is False


def test_post_publication_sync_failure_is_unknown_and_never_replayed(
    document, monkeypatch, capsys
):
    workspace, source, request = document
    (source.parent / "changes.json").write_text(json.dumps(request))
    fsync = os.fsync
    count = 0

    def fail_directory_sync(descriptor):
        nonlocal count
        count += 1
        if count == 2:
            raise OSError("directory sync failed")
        return fsync(descriptor)

    monkeypatch.setattr(os, "fsync", fail_directory_sync)
    assert (
        presentation.main(
            ["edit", "--workspace", str(workspace.root), "--request", "changes.json"]
        )
        == 2
    )
    result = json.loads(capsys.readouterr().out)
    assert result["writePublished"] is None and result["error"] == "PUBLICATION_UNKNOWN"
    assert source.with_name("edited.pptx").exists()
