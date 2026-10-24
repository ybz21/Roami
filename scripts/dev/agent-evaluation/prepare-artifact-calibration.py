"""Prepare synthetic calibration evidence without claiming human review."""

import argparse
import asyncio
import base64
import html
import json
from pathlib import Path

from roami_agent.improvement.artifact_evidence import PROGRAM
from roami_agent.improvement.containers import ContainerExecutor
from roami_agent.improvement.contracts import DIMENSIONS, digest
from roami_agent.improvement.executors import private_json
from roami_agent.server.configuration import Settings

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--output-root", type=Path, required=True)
parser.add_argument("--review-root", type=Path, required=True)
parser.add_argument("--renderer-image", required=True)
args = parser.parse_args()
ROOT = args.output_root.resolve()
PUBLIC = args.review_root.resolve()
ROOT.mkdir(mode=0o700, parents=True, exist_ok=False)
PUBLIC.mkdir(parents=True, exist_ok=False)
DECK = r"""import io,base64,json
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
s=json.loads(input());r=Presentation();r.slide_width=Inches(10);r.slide_height=Inches(5.625)
for title,body in s['pages']:
 slide=r.slides.add_slide(r.slide_layouts[6]);background=slide.background.fill;background.solid();background.fore_color.rgb=RGBColor(255,255,255)
 if not s['bad']:
  for y,text,size in [(0.4,title,30),(1.4,body,23),(4.95,s['source'],12)]:
   box=slide.shapes.add_textbox(Inches(.6),Inches(y),Inches(8.8),Inches(3 if y==1.4 else .6));tf=box.text_frame;tf.word_wrap=True
   for i,line in enumerate(text.split('\n')):
    p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Liberation Sans';p.font.size=Pt(size);p.font.color.rgb=RGBColor(18,32,48)
b=io.BytesIO();r.save(b);print(base64.b64encode(b.getvalue()).decode())"""
SCENARIOS = {
    "presentation-sales": (
        "presentation",
        "Create two readable slides. Show Q1=20, Q2=35 and their total=55 units. Cite synthetic-sales.csv.",
        {"facts": "synthetic-sales.csv: Q1=20; Q2=35; total=55"},
        [
            ["Quarterly sales", "Q1: 20 units\nQ2: 35 units"],
            ["Total sales", "20 + 35 = 55 units\nQ2 exceeds Q1 by 15 units."],
        ],
    ),
    "presentation-release": (
        "presentation",
        "Create two readable slides showing owner Ada, release date 2026-11-02 and rollback target 60 seconds. Cite synthetic-release.json.",
        {
            "facts": "synthetic-release.json: owner Ada; release 2026-11-02; rollback target 60 seconds"
        },
        [
            ["Release plan", "Owner: Ada\nRelease: 2026-11-02"],
            [
                "Recovery plan",
                "Rollback target: 60 seconds\nRestore the previous version if the health check fails.",
            ],
        ],
    ),
    "plan-import": (
        "plan",
        "Design an internal CSV import service. Specify POST /v1/import, requestId idempotency, validation errors and recovery after timeout. Include a diagram and acceptance criteria.",
        {
            "facts": "Internal offline design exercise; CSV columns name,value. requestId identifies one import. Validation errors must not write rows."
        },
        "# CSV import design\n\n## Product and story\nAn operator imports name,value rows once; retries after a timeout preserve the same requestId.\n\n## Interface and storage\nPOST /v1/import accepts JSON {requestId:string, rows:[{name:string,value:number}]}. A transaction stores an import record keyed by requestId and its rows. Identical retries return the stored result. Reusing the key for different content returns 409. GET /v1/import/{requestId} reads status.\n\n## Errors and recovery\nValidate all rows before the transaction. Return 400 with row numbers for missing columns or nonnumeric values, without writes. If the response times out, read status and retry with the same key. Roll back a failed transaction.\n\n## Acceptance\nValid rows are saved once. Invalid rows cause no partial writes. A timeout followed by retry adds no duplicate rows.\n\nSource: synthetic-import-requirements.\n",
    ),
    "plan-backup": (
        "plan",
        "Design nightly encrypted backups retained for 7 days. Describe the scheduler, backup storage, checksum verification, failed-backup alert and restore check. Include a diagram and acceptance criteria.",
        {
            "facts": "Synthetic backup requirements: nightly snapshot; encryption; retention 7 days; checksum; failure alert; restore verification."
        },
        "# Backup design\n\n## Product and story\nThe operator receives one encrypted snapshot each night and can restore a checksum-verified copy.\n\n## Components and interface\nA scheduler creates a dated backup job. The worker creates a consistent snapshot, encrypts it, writes a temporary object, verifies its checksum, then atomically publishes the object and manifest. POST /v1/backups creates a job; GET /v1/backups/{id} returns its state and checksum.\n\n## Retention and recovery\nRetain completed snapshots for 7 days. A failed upload deletes only its temporary object, emits a failure alert, and keeps the last good snapshot. Restore into an isolated directory, verify the checksum and run a sample read before switching the restored database.\n\n## Acceptance\nThe manifest is published only after verification. Interrupted uploads preserve the previous backup. A restore check reproduces the verified data.\n\nSource: synthetic-backup-requirements.\n",
    ),
    "web-counter": (
        "web",
        "Build an offline counter. Initially show 0, Increase adds one and Reset returns to zero. Buttons must work with keyboard and be readable on phone and desktop.",
        {
            "facts": "Counter initial value 0; Increase +1; Reset 0; no external network."
        },
        '<h1>Counter</h1><p id="value" aria-live="polite">0</p><button id="inc" onclick="value.textContent=Number(value.textContent)+1">Increase</button><button onclick="value.textContent=0">Reset</button>',
    ),
    "web-filter": (
        "web",
        "Build an offline searchable list containing Ada, Ben and Cy. Label the search field and show the matching count. It must fit phone and desktop.",
        {
            "facts": "Provided names: Ada, Ben, Cy. Case-insensitive name filtering. No external network."
        },
        '<h1>People</h1><label for="search">Search names</label><input id="search" oninput="render()"><p id="count" aria-live="polite">3 results</p><ul id="names"><li>Ada</li><li>Ben</li><li>Cy</li></ul><script>function render(){const a=["Ada","Ben","Cy"].filter(x=>x.toLowerCase().includes(search.value.toLowerCase()));names.replaceChildren(...a.map(x=>{const e=document.createElement("li");e.textContent=x;return e}));count.textContent=a.length+" results"}</script>',
    ),
}
SVG = (
    '<svg xmlns="http://www.w3.org/2000/svg" width="900" height="220" viewBox="0 0 900 220"><rect width="900" height="220" fill="white"/>'
    + "".join(
        f'<rect x="{x}" y="65" width="220" height="80" fill="#edf3fa" stroke="#334155"/><text x="{x + 110}" y="110" text-anchor="middle" font-size="22" font-family="sans-serif">{t}</text>'
        for x, t in [
            (20, "Client / Scheduler"),
            (335, "Validated worker"),
            (650, "Transactional store"),
        ]
    )
    + '<path d="M240 105H335M555 105H650" stroke="#334155" stroke-width="3"/></svg>'
)


async def main():
    executor = ContainerExecutor(Settings(), image=args.renderer_image)
    await executor.preflight()
    anchors = []
    for name, (family, instruction, grounding, content) in SCENARIOS.items():
        for bad in (False, True):
            ident = name + ("-incomplete" if bad else "-complete")
            folder = ROOT / ident
            folder.mkdir(mode=0o700, exist_ok=True)
            workspace = folder / "workspace"
            workspace.mkdir(mode=0o700, exist_ok=True)
            if family == "presentation":
                raw = await executor.sandbox(
                    workspace,
                    DECK,
                    input=json.dumps(
                        {
                            "pages": content,
                            "source": grounding["facts"].split(":")[0],
                            "bad": bad,
                        }
                    ).encode(),
                )
                (workspace / "presentation.pptx").write_bytes(base64.b64decode(raw))
                files = ["presentation.pptx"]
            elif family == "plan":
                (workspace / "plan.md").write_text(
                    "TODO: unspecified service. No interface, recovery or acceptance criteria."
                    if bad
                    else content
                )
                (workspace / "architecture.svg").write_text(
                    '<svg xmlns="http://www.w3.org/2000/svg" width="900" height="220"><rect width="900" height="220" fill="white"/></svg>'
                    if bad
                    else SVG
                )
                files = ["plan.md", "architecture.svg"]
            else:
                head = '<meta name="viewport" content="width=device-width,initial-scale=1"><style>body{font:20px/1.6 system-ui;margin:24px;max-width:850px;color:#172033;background:white}button,input{font:inherit;padding:12px;margin:4px;max-width:100%;box-sizing:border-box}#value{font-size:40px}input{display:block}</style>'
                (workspace / "index.html").write_text(
                    "<!doctype html>" + head + ("" if bad else content)
                )
                files = ["index.html"]
            rendered = json.loads(
                await executor.sandbox(
                    workspace,
                    PROGRAM,
                    input=json.dumps(
                        {"id": ident, "files": files, "web": family == "web"}
                    ).encode(),
                    timeout=180,
                )
            )
            import hashlib

            for image in rendered["images"]:
                image["sha256"] = hashlib.sha256(
                    base64.b64decode(image["data"])
                ).hexdigest()
            value = {
                "schema_version": 1,
                "renderer_image": executor.image_id,
                "instruction": instruction,
                "grounding": grounding,
                "artifacts": [
                    {
                        "file": f,
                        "sha256": hashlib.sha256(
                            (workspace / f).read_bytes()
                        ).hexdigest(),
                    }
                    for f in files
                ],
                **rendered,
            }
            pack = {**value, "content_hash": digest(value)}
            private_json(folder / "evidence.json", pack)
            anchor = {
                "id": ident,
                "family": family,
                "human_reference": False,
                "human_rater": "",
                "expected": dict.fromkeys(DIMENSIONS, 1 if bad else 3),
                "evidence_file": str(folder / "evidence.json"),
                "evidence_hash": pack["content_hash"],
            }
            anchors.append(anchor)
            print(
                json.dumps({"anchor": ident, "images": len(pack["images"])}), flush=True
            )
            item = PUBLIC / ident
            item.mkdir(exist_ok=True)
            for index, img in enumerate(pack["images"]):
                (item / (str(index) + ".jpg")).write_bytes(
                    base64.b64decode(img["data"])
                )
            (item / "pack.json").write_text(
                json.dumps(
                    {k: v for k, v in pack.items() if k != "images"},
                    ensure_ascii=False,
                    indent=2,
                )
            )
    private_json(ROOT / "proposed-anchors.json", anchors)
    rows = []
    for a in anchors:
        pack = json.loads(Path(a["evidence_file"]).read_text())
        ident = a["id"]
        score = 1 if ident.endswith("incomplete") else 3
        rows.append(
            '<section id="'
            + ident
            + '"><h2>'
            + html.escape(ident)
            + "</h2><p>"
            + html.escape(pack["instruction"])
            + "</p><pre>"
            + html.escape(json.dumps(pack["grounding"], ensure_ascii=False, indent=2))
            + "</pre><p>建议参考分："
            + str(score)
            + " / 3（四个维度均待人工核对）</p>"
            + "".join(
                '<img loading="lazy" src="'
                + ident
                + "/"
                + str(i)
                + '.jpg" alt="'
                + html.escape(im["label"])
                + '">'
                for i, im in enumerate(pack["images"])
            )
            + "<details><summary>查看正文</summary><pre>"
            + html.escape("\n\n".join(t["text"] for t in pack["texts"]))
            + "</pre></details></section>"
        )
    page = (
        '<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>产物评分校准 · 12项待确认</title><style>body{max-width:1100px;margin:auto;padding:20px;font:16px/1.6 system-ui;background:#0d1117;color:#e6edf3}a{color:#58a6ff}section{border-top:1px solid #30363d;padding-block:24px}img{display:block;max-width:100%;height:auto;margin:12px 0}pre{white-space:pre-wrap;overflow-wrap:anywhere}nav{display:flex;flex-wrap:wrap;gap:12px}</style><h1>产物评分校准 · 待人工确认</h1><p>12个独立合成样本，分别覆盖演示文稿、方案和网页的完整与缺陷产物。以下分数由Codex提出，尚不是人工参考分；不签发校准证书、不放行正式评测。</p><p>维度：正确性、完整性、可用性、依据充分性。0=不可用；1=严重缺陷；2=部分完成；3=完整满足。每个完整样本建议四项3分；每个缺陷样本建议四项1分。可逐项修改。</p><nav>'
        + "".join('<a href="#' + a["id"] + '">' + a["id"] + "</a>" for a in anchors)
        + "</nav>"
        + "".join(rows)
        + "</html>"
    )
    (PUBLIC / "index.html").write_text(page)
    private_json(
        ROOT / "manifest.json",
        {
            "renderer_image": executor.image_id,
            "anchors": anchors,
            "human_confirmed": False,
        },
    )
    print("review-ready", flush=True)


asyncio.run(main())
